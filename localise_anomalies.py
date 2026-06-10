#!/usr/bin/env python3
"""Localise les couleurs et tailles anormales."""

from docx import Document
from docx.oxml.ns import qn

doc = Document('/home/user/tts-docs/LA_PAIX_ON_PEUT_L_EVITER_CORRIGE.docx')
paras = doc.paragraphs

# Couleurs à signaler
BAD_COLORS  = {'8B0000', '943634', '555555', '404040', '4A4A4A', 'D99594', '1A1A1A'}
# Tailles anormales en EMU
BAD_SIZES   = {355600, 152400}   # 28pt, 12pt
BAD_FONTS   = {'Segoe UI Symbol'}

print("=== POLICES ANORMALES ===")
for i, p in enumerate(paras):
    for run in p.runs:
        if run.font.name in BAD_FONTS:
            print(f"  [{i:4d}] font={run.font.name!r} | {p.text.strip()[:60]!r}")

print("\n=== TAILLES ANORMALES ===")
for i, p in enumerate(paras):
    for run in p.runs:
        if run.font.size in BAD_SIZES:
            pt = run.font.size / 12700
            print(f"  [{i:4d}] {pt}pt | style={p.style.name!r} | {p.text.strip()[:60]!r}")

print("\n=== COULEURS ANORMALES ===")
for i, p in enumerate(paras):
    for run in p.runs:
        if not run.text.strip():
            continue
        try:
            clr = str(run.font.color.rgb) if run.font.color and run.font.color.type else None
        except:
            clr = None
        if clr and clr.upper() in BAD_COLORS:
            print(f"  [{i:4d}] #{clr} | style={p.style.name!r} | {p.text.strip()[:60]!r}")

# Couleur des tableaux
print("\n=== COULEURS DANS LES TABLEAUX ===")
for t_idx, table in enumerate(doc.tables):
    for row in table.rows:
        for cell in row.cells:
            for cp in cell.paragraphs:
                for run in cp.runs:
                    if not run.text.strip():
                        continue
                    try:
                        clr = str(run.font.color.rgb) if run.font.color and run.font.color.type else None
                    except:
                        clr = None
                    if clr and clr.upper() in BAD_COLORS:
                        print(f"  Table[{t_idx}] #{clr} | {cp.text.strip()[:50]!r}")
