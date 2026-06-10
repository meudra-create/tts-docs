#!/usr/bin/env python3
"""
CORRECTION GLOBALE — LA PAIX, ON PEUT L'ÉVITER (BEN-H2O, 2026)
Applique l'intégralité des règles format-livre-kemi + prompt-editeur-kdp
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CHEMIN = '/home/user/tts-docs/LA_PAIX_ON_PEUT_L_EVITER_CORRIGE.docx'
OUTPUT = CHEMIN

doc = Document(CHEMIN)
paras = doc.paragraphs
print(f"Document : {len(paras)} paragraphes, {len(doc.tables)} tableaux\n")

# ═══════════════════════════════════════════════════════════════
# MODULE A — format-livre-kemi
# ═══════════════════════════════════════════════════════════════

# ─── A1 : Titres chapitres (Heading 2) ────────────────────────
print("A1 — Titres chapitres (Heading 2)...")
h2_count = 0
for p in paras:
    if p.style.name != 'Heading 2':
        continue
    for run in p.runs:
        run.font.name  = 'Garamond'
        run.font.size  = 177800       # 14pt
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        run.font.bold  = False
    pf = p.paragraph_format
    pf.page_break_before = True
    pf.keep_with_next    = True
    pf.space_before      = Pt(0)
    pf.space_after       = Pt(12)
    h2_count += 1
print(f"  ✓ {h2_count} chapitres Heading 2")

# ─── A2 : Citations «...» ─────────────────────────────────────
print("A2 — Citations...")
cit_count = 0
for p in paras:
    if not p.text.strip().startswith('«'):
        continue
    # Police & taille
    for run in p.runs:
        run.font.size   = 133350      # 10.5pt
        run.font.italic = True
    # Espacement & retrait
    pf = p.paragraph_format
    pf.space_before  = Pt(8)
    pf.space_after   = Pt(2)
    pf.left_indent   = 432000        # 1.2cm
    pf.right_indent  = 288000        # 0.8cm
    pf.first_line_indent = 0         # pas de retrait sur citations

    # Barre bordeaux gauche
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p._p.insert(0, pPr)
    old = pPr.find(qn('w:pBdr'))
    if old is not None:
        pPr.remove(old)
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'),   'single')
    left.set(qn('w:sz'),    '32')    # 4pt
    left.set(qn('w:space'), '14')
    left.set(qn('w:color'), '7B1E1E')
    pBdr.append(left)
    pPr.append(pBdr)
    cit_count += 1
print(f"  ✓ {cit_count} citations")

# ─── A3 : Attributions — Auteur ──────────────────────────────
print("A3 — Attributions...")
attr_count = 0
for p in paras:
    txt = p.text.strip()
    if not (txt.startswith('—') or txt.startswith('–')):
        continue
    if p.style.name in ('Heading 1', 'Heading 2'):
        continue
    for run in p.runs:
        run.font.size   = 120650     # 9.5pt
        run.font.italic = True
    pf = p.paragraph_format
    pf.space_before  = Pt(0)
    pf.space_after   = Pt(10)
    pf.left_indent   = 432000
    pf.right_indent  = 288000
    pf.first_line_indent = 0
    # S'assurer qu'il n'y a PAS de barre bordeaux sur les attributions
    pPr = p._p.find(qn('w:pPr'))
    if pPr is not None:
        old = pPr.find(qn('w:pBdr'))
        if old is not None:
            pPr.remove(old)
    attr_count += 1
print(f"  ✓ {attr_count} attributions")

# ─── A4 : Étoiles séparatrices ───────────────────────────────
print("A4 — Étoiles...")
star_count = 0
for p in paras:
    if '★' not in p.text:
        continue
    # Normaliser texte
    for run in p.runs:
        run.text = ''
    if not p.runs:
        p.add_run('★ ★ ★')
    else:
        p.runs[0].text = '★ ★ ★'
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x7B, 0x1E, 0x1E)
        run.font.bold = False
        run.font.italic = False
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after  = Pt(10)
    pf.first_line_indent = 0
    star_count += 1
print(f"  ✓ {star_count} étoiles")

# ─── A5 : Sommaire ───────────────────────────────────────────
print("A5 — Sommaire [47–155]...")
som_partie = som_chap = 0
KEYWORDS_PARTIE = ['PARTIE', 'ÉPILOGUE', 'EPILOGUE', 'PROLOGUE',
                   'INTRODUCTION', 'AVANT-PROPOS', 'CONCLUSION']
for i, p in enumerate(paras):
    if i < 47 or i > 155:
        continue
    txt = p.text.strip()
    if not txt:
        continue
    upper = txt.upper()
    pf = p.paragraph_format
    if any(kw in upper for kw in KEYWORDS_PARTIE):
        pf.left_indent  = 0
        pf.space_before = Pt(12)
        pf.space_after  = Pt(4)
        pf.first_line_indent = 0
        for run in p.runs:
            run.font.bold = True
        som_partie += 1
    elif 'Chapitre' in txt or 'CHAPITRE' in upper:
        pf.left_indent  = Cm(1.5)
        pf.space_before = Pt(3)
        pf.space_after  = 0
        pf.first_line_indent = 0
        for run in p.runs:
            run.font.bold = False
        som_chap += 1
print(f"  ✓ {som_partie} PARTIE, {som_chap} chapitres dans le sommaire")

# ─── A6 : Tableaux ───────────────────────────────────────────
print("A6 — Tableaux...")
for table in doc.tables:
    # En-tête : sp_before=14pt
    header_para = table.rows[0].cells[0].paragraphs[0]
    pPr = header_para._p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        header_para._p.insert(0, pPr)
    spacing = pPr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        pPr.append(spacing)
    spacing.set(qn('w:before'), '280')   # 14pt
    # Normaliser fond en-tête #1A1A2E
    tc = table.rows[0].cells[0]._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is not None:
        shd = tcPr.find(qn('w:shd'))
        if shd is not None:
            shd.set(qn('w:fill'), '1A1A2E')
    # Dernière ligne : sp_after=8pt
    last_para = table.rows[-1].cells[0].paragraphs[-1]
    pPr2 = last_para._p.find(qn('w:pPr'))
    if pPr2 is None:
        pPr2 = OxmlElement('w:pPr')
        last_para._p.insert(0, pPr2)
    sp2 = pPr2.find(qn('w:spacing'))
    if sp2 is None:
        sp2 = OxmlElement('w:spacing')
        pPr2.append(sp2)
    sp2.set(qn('w:after'), '160')       # 8pt
print(f"  ✓ {len(doc.tables)} tableaux")

# ═══════════════════════════════════════════════════════════════
# MODULE B — prompt-editeur-kdp
# ═══════════════════════════════════════════════════════════════

# ─── B1 : Césure ─────────────────────────────────────────────
print("\nB1 — Césure automatique...")
settings = doc.settings.element
for tag, val in [('w:autoHyphenation', '1'),
                 ('w:hyphenationZone', '720'),
                 ('w:doNotHyphenateCaps', '1')]:
    el = settings.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        settings.insert(0, el)
    el.set(qn('w:val'), val)
print("  ✓ autoHyphenation activée")

# ─── B2 : Titres de parties (Heading 1) ──────────────────────
print("B2 — Titres de parties (Heading 1)...")
h1_count = 0
for p in paras:
    if p.style.name != 'Heading 1':
        continue
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.page_break_before = True
    pf.keep_with_next    = True
    pf.space_before      = Pt(36)
    pf.space_after       = Pt(12)
    for run in p.runs:
        run.font.name  = 'Garamond'
        run.font.size  = 241300       # 19pt
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        run.font.bold  = True
    # Bordure gold
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p._p.insert(0, pPr)
    old_bdr = pPr.find(qn('w:pBdr'))
    if old_bdr is None:
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        for k, v in [('w:val','single'),('w:sz','8'),
                     ('w:space','8'),('w:color','C8A951')]:
            bottom.set(qn(k), v)
        pBdr.append(bottom)
        pPr.append(pBdr)
    h1_count += 1
print(f"  ✓ {h1_count} PARTIES Heading 1")

# ─── B3 : Sections liminaires ────────────────────────────────
print("B3 — Sections liminaires...")
def add_gold_bottom(p):
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); p._p.insert(0, pPr)
    if pPr.find(qn('w:pBdr')) is None:
        pBdr = OxmlElement('w:pBdr')
        bot  = OxmlElement('w:bottom')
        for k, v in [('w:val','single'),('w:sz','8'),
                     ('w:space','8'),('w:color','C8A951')]:
            bot.set(qn(k), v)
        pBdr.append(bot); pPr.append(pBdr)

# Détection par texte (robuste si indices changent)
LIMINAIRE_KEYWORDS = ['AVANT-PROPOS', 'AVANT PROPOS', 'SOMMAIRE',
                      'PROLOGUE', 'INTRODUCTION', 'ÉPILOGUE', 'EPILOGUE']
lim_count = 0
for i, p in enumerate(paras):
    txt = p.text.strip()
    if not txt or p.style.name in ('Heading 1', 'Heading 2'):
        continue
    upper = txt.upper()
    if any(kw in upper for kw in LIMINAIRE_KEYWORDS) and len(txt) < 80:
        # Pas dans le sommaire
        if 47 <= i <= 155:
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.page_break_before = True
        pf.space_before      = Pt(24)
        pf.space_after       = Pt(18)
        pf.keep_with_next    = True
        pf.first_line_indent = 0
        for run in p.runs:
            run.font.name  = 'Garamond'
            run.font.size  = 215900   # 17pt
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
            run.font.bold  = True
        add_gold_bottom(p)
        lim_count += 1
print(f"  ✓ {lim_count} sections liminaires")

# ─── B4 : Page de titre ──────────────────────────────────────
print("B4 — Page de titre...")
# Détection robuste par contenu
TITRE_CONFIGS = {
    'Alliance des États du Sahel': (139700, False, True,  RGBColor(0x1A,0x1A,0x2E), Pt(48), Pt(6)),
    'LA PAIX,':                    (406400, True,  False, RGBColor(0x1A,0x1A,0x2E), Pt(24), Pt(0)),
    "ON PEUT L'ÉVITER":            (406400, True,  False, RGBColor(0x1A,0x1A,0x2E), Pt(0),  Pt(18)),
    'Géopolitique de la résistance africaine': (165100, False, True, RGBColor(0x1A,0x1A,0x2E), Pt(6), Pt(60)),
    'BEN':                         (190500, True,  False, RGBColor(0x1A,0x1A,0x2E), Pt(60), Pt(4)),
    '— 2026 —':                    (139700, False, False, RGBColor(0x7B,0x1E,0x1E), Pt(2),  Pt(2)),
}
titre_count = 0
for p in paras[:30]:
    txt = p.text.strip()
    for key, (size, bold, italic, color, sp_b, sp_a) in TITRE_CONFIGS.items():
        if txt.startswith(key):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before      = sp_b
            pf.space_after       = sp_a
            pf.first_line_indent = 0
            for run in p.runs:
                run.font.name   = 'Garamond'
                run.font.size   = size
                run.font.bold   = bold
                run.font.italic = italic
                run.font.color.rgb = color
            titre_count += 1
            break
# Sous-titre : bordure gold
for p in paras[:30]:
    if p.text.strip().startswith('Géopolitique'):
        add_gold_bottom(p)
# Épigraphes : page dédiée
for p in paras[20:40]:
    if p.text.strip().startswith('«') and paras.index(p) < 40:
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(48)
        break
print(f"  ✓ {titre_count} éléments page de titre")

# ─── B5 : Corps de texte — interligne + justification ────────
print("B5 — Corps de texte...")
body_count = 0
SKIP_STYLES = {'Heading 1','Heading 2','Heading 3','Title','Subtitle',
               'TOC 1','TOC 2','TOC 3','Caption','List Paragraph'}
for p in paras:
    if p.style.name in SKIP_STYLES:
        continue
    txt = p.text.strip()
    if not txt or len(txt) < 20:
        continue
    if txt.startswith(('«','—','–')) or '★' in txt:
        continue
    # Appliquer seulement sur styles corps
    if p.style.name in ('Normal','Body Text','') or p.style.name.startswith('Normal'):
        pf = p.paragraph_format
        # Justification
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Interligne 1.3×
        pf.line_spacing      = 1.3
        pf.widow_control     = True
        # Espacement après 6pt
        pf.space_after       = Pt(6)
        # Contrôle veuves/orphelines via XML
        pPr = p._p.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p._p.insert(0, pPr)
        wc = pPr.find(qn('w:widowControl'))
        if wc is None:
            wc = OxmlElement('w:widowControl')
            pPr.insert(0, wc)
        wc.set(qn('w:val'), '1')
        body_count += 1
print(f"  ✓ {body_count} paragraphes corps de texte")

# ─── B6 : Retrait première ligne 0.5cm ──────────────────────
print("B6 — Retrait première ligne...")
FIRST_LINE = 180000   # 0.5cm en EMU
skip_next = False
indent_count = 0
for i, p in enumerate(paras):
    if p.style.name in ('Heading 1','Heading 2'):
        skip_next = True
        continue
    txt = p.text.strip()
    if not txt:
        continue
    if 47 <= i <= 155:         # sommaire → pas de retrait
        skip_next = False
        continue
    if p.style.name in SKIP_STYLES:
        skip_next = False
        continue
    if txt.startswith(('«','—','–')) or '★' in txt:
        skip_next = False
        continue
    if any(kw in txt.upper() for kw in LIMINAIRE_KEYWORDS) and len(txt) < 80:
        skip_next = False
        continue
    if len(txt) < 40:
        skip_next = False
        continue
    # 1er § après un titre → pas de retrait (règle typographique)
    if skip_next:
        p.paragraph_format.first_line_indent = 0
        skip_next = False
        continue
    if p.style.name in ('Normal','Body Text','') or p.style.name.startswith('Normal'):
        p.paragraph_format.first_line_indent = FIRST_LINE
        indent_count += 1
print(f"  ✓ {indent_count} paragraphes avec retrait 0.5cm")

# ═══════════════════════════════════════════════════════════════
# SAUVEGARDE
# ═══════════════════════════════════════════════════════════════
print("\n--- SAUVEGARDE ---")
doc.save(OUTPUT)
print(f"✓ {OUTPUT}")

print("""
╔══════════════════════════════════════════════════════════════╗
║          CORRECTION GLOBALE — RAPPORT FINAL                 ║
╠══════════════════════════════════════════════════════════════╣
║  MODULE A — format-livre-kemi                               ║
║  A1 ✓ Titres chapitres : Garamond 14pt #1A1A2E              ║
║  A2 ✓ Citations : barre bordeaux sz=32, retrait G/D          ║
║  A3 ✓ Attributions : 9.5pt italique, retrait G/D            ║
║  A4 ✓ Étoiles : ★ ★ ★ bordeaux centrées                     ║
║  A5 ✓ Sommaire : PARTIE bold/0, Chapitre indent 1.5cm       ║
║  A6 ✓ Tableaux : sp_before 14pt, sp_after 8pt, fond navy    ║
║                                                              ║
║  MODULE B — prompt-editeur-kdp                              ║
║  B1 ✓ Césure automatique (zone 720 twips)                   ║
║  B2 ✓ PARTIES centrées, Garamond 19pt, saut de page         ║
║  B3 ✓ Liminaires : 17pt gras centré bordure gold            ║
║  B4 ✓ Page de titre : 32pt gras, sous-titre gold            ║
║  B5 ✓ Corps : justifié, interligne 1.3×, sp_after 6pt       ║
║  B6 ✓ Retrait première ligne 0.5cm (sauf 1er § après titre) ║
║                                                              ║
║  FORMAT KDP 6×9 pouces maintenu                             ║
╚══════════════════════════════════════════════════════════════╝
""")
