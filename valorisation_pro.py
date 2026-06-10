#!/usr/bin/env python3
"""
Valorisation professionnelle de la présentation
LA PAIX, ON PEUT L'ÉVITER (BEN-H2O, 2026)

Ajoute les éléments d'une édition professionnelle :
1. Faux-titre (page de demi-titre, tradition éditoriale)
2. Fleuron décoratif ❦ sur la page de titre
3. Page de copyright / mentions légales
4. En-têtes courants alternés (titre recto / auteur verso)
5. Numérotation des pages en pied (champ PAGE)
6. Première page sans en-tête ni numéro
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CHEMIN = '/home/user/tts-docs/LA_PAIX_ON_PEUT_L_EVITER_CORRIGE.docx'

NAVY     = RGBColor(0x1A, 0x1A, 0x2E)
BORDEAUX = RGBColor(0x7B, 0x1E, 0x1E)

doc = Document(CHEMIN)
paras = doc.paragraphs
print(f"Document : {len(paras)} paragraphes\n")

# Garde-fou : ne pas dupliquer si déjà appliqué
deja_fait = any('Tous droits réservés' in p.text for p in paras[:60])
if deja_fait:
    print("⚠ Valorisation déjà appliquée — abandon pour éviter les doublons.")
    raise SystemExit(0)

# ────────────────────────────────────────────────────────────
# 1. FAUX-TITRE (demi-titre) — première page, titre seul
# ────────────────────────────────────────────────────────────
print("1 — Faux-titre...")
# paras[2] = 'Alliance des États du Sahel' (début de la page de titre)
premiere = paras[2]
fx = premiere.insert_paragraph_before("LA PAIX, ON PEUT L'ÉVITER")
fx.alignment = WD_ALIGN_PARAGRAPH.CENTER
fx.paragraph_format.space_before = Pt(160)
fx.paragraph_format.space_after  = Pt(0)
for run in fx.runs:
    run.font.name      = 'Garamond'
    run.font.size      = 203200       # 16pt
    run.font.color.rgb = NAVY
    run.font.bold      = False
    # petites capitales — élégance éditoriale
    rPr = run._r.find(qn('w:rPr'))
    sc  = OxmlElement('w:smallCaps')
    sc.set(qn('w:val'), '1')
    rPr.append(sc)

# La page de titre commence sur une nouvelle page
premiere.paragraph_format.page_break_before = True
print("  ✓ Faux-titre en petites capitales, Garamond 16pt")

# ────────────────────────────────────────────────────────────
# 2. FLEURON ❦ sur la page de titre (après l'année)
# ────────────────────────────────────────────────────────────
print("2 — Fleuron...")
paras = doc.paragraphs   # rafraîchir après insertion
year_idx = None
for i, p in enumerate(paras[:30]):
    if p.text.strip() == '— 2026 —':
        year_idx = i
        break

if year_idx is not None:
    suivant = paras[year_idx + 1]
    fl = suivant.insert_paragraph_before('❦')
    fl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fl.paragraph_format.space_before = Pt(20)
    for run in fl.runs:
        run.font.name      = 'Garamond'
        run.font.size      = 177800   # 14pt
        run.font.color.rgb = BORDEAUX
    print("  ✓ Fleuron ❦ bordeaux sous l'année")
else:
    print("  ⚠ Année non trouvée — fleuron non inséré")

# ────────────────────────────────────────────────────────────
# 3. PAGE DE COPYRIGHT / MENTIONS LÉGALES
# ────────────────────────────────────────────────────────────
print("3 — Page de copyright...")
paras = doc.paragraphs
# La page suivant la page de titre = première épigraphe «...»
epigraphe = None
for p in paras[:45]:
    if p.text.strip().startswith('«'):
        epigraphe = p
        break

MENTIONS = [
    ('LA PAIX, ON PEUT L’ÉVITER', True,  Pt(260), Pt(2)),
    ('Géopolitique de la résistance africaine', False, Pt(0), Pt(14)),
    ('© 2026 BEN–H2O — Tous droits réservés.', False, Pt(0), Pt(8)),
    ('Aucune partie de cet ouvrage ne peut être reproduite, stockée ou '
     'transmise sous quelque forme ou par quelque moyen que ce soit — '
     'électronique, mécanique, photocopie, enregistrement ou autre — '
     'sans l’autorisation écrite préalable de l’auteur.', False, Pt(0), Pt(8)),
    ('Première édition : 2026', False, Pt(0), Pt(2)),
    ('Format : 15,24 × 22,86 cm (6 × 9 po)', False, Pt(0), Pt(2)),
    ('ISBN : [à compléter]', False, Pt(0), Pt(2)),
    ('Dépôt légal : [à compléter]', False, Pt(0), Pt(2)),
    ('Publié à compte d’auteur via Kindle Direct Publishing.', False, Pt(0), Pt(0)),
]

if epigraphe is not None:
    first_line = True
    for texte, bold, sp_b, sp_a in MENTIONS:
        cp = epigraphe.insert_paragraph_before(texte)
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = cp.paragraph_format
        pf.space_before      = sp_b
        pf.space_after       = sp_a
        pf.first_line_indent = 0
        pf.line_spacing      = 1.15
        if first_line:
            pf.page_break_before = True
            first_line = False
        for run in cp.runs:
            run.font.name      = 'Garamond'
            run.font.size      = 107950   # 8.5pt
            run.font.bold      = bold
            run.font.color.rgb = NAVY if bold else RGBColor(0x40, 0x40, 0x40)
    print(f"  ✓ {len(MENTIONS)} mentions légales (Garamond 8.5pt, bas de page)")
else:
    print("  ⚠ Épigraphe non trouvée — copyright non inséré")

# ────────────────────────────────────────────────────────────
# 4. EN-TÊTES COURANTS ALTERNÉS + 5. NUMÉROTATION DES PAGES
# ────────────────────────────────────────────────────────────
print("4/5 — En-têtes courants et numérotation...")

doc.settings.odd_and_even_pages_header_footer = True

def style_header_par(par, text, italic=True):
    par.text = ''
    run = par.add_run(text)
    run.font.name      = 'Garamond'
    run.font.size      = 114300       # 9pt
    run.font.italic    = italic
    run.font.color.rgb = NAVY
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)

def add_page_field(par):
    par.text = ''
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), ' PAGE ')
    r  = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf  = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), 'Garamond')
    rf.set(qn('w:hAnsi'), 'Garamond')
    sz  = OxmlElement('w:sz')
    sz.set(qn('w:val'), '18')         # 9pt
    rPr.append(rf)
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fld.append(r)
    par._p.append(fld)

for section in doc.sections:
    section.different_first_page_header_footer = True

    # Recto (pages impaires) : titre du livre
    h = section.header
    h.is_linked_to_previous = False
    style_header_par(h.paragraphs[0], "LA PAIX, ON PEUT L'ÉVITER")

    # Verso (pages paires) : auteur
    he = section.even_page_header
    he.is_linked_to_previous = False
    style_header_par(he.paragraphs[0], 'BEN–H2O')

    # Pieds de page : numéro centré
    f = section.footer
    f.is_linked_to_previous = False
    add_page_field(f.paragraphs[0])

    fe = section.even_page_footer
    fe.is_linked_to_previous = False
    add_page_field(fe.paragraphs[0])

    # Première page (faux-titre) : ni en-tête ni numéro
    section.first_page_header.is_linked_to_previous = False
    section.first_page_header.paragraphs[0].text = ''
    section.first_page_footer.is_linked_to_previous = False
    section.first_page_footer.paragraphs[0].text = ''

print("  ✓ Recto : titre · Verso : auteur · Numéros centrés · 1re page vierge")

# ────────────────────────────────────────────────────────────
# SAUVEGARDE
# ────────────────────────────────────────────────────────────
doc.save(CHEMIN)
paras = doc.paragraphs
print(f"\n✓ Sauvegardé : {CHEMIN} ({len(paras)} paragraphes)")
