#!/usr/bin/env python3
"""
kdp_formatter.py — Formateur KDP professionnel
LA PAIX, ON PEUT L'ÉVITER (BEN-H2O, 2026)

Usage :
    python3 kdp_formatter.py
    python3 kdp_formatter.py --input mon_fichier.docx --output sortie.docx
"""

import sys
import argparse
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Palette ───────────────────────────────────────────────────
NAVY     = RGBColor(0x1A, 0x1A, 0x2E)
BORDEAUX = RGBColor(0x7B, 0x1E, 0x1E)
GOLD     = 'C8A951'

# ── Tailles (EMU) ─────────────────────────────────────────────
SZ_TITRE      = 406400   # 32pt  — page de titre
SZ_PARTIE     = 241300   # 19pt  — Heading 1
SZ_LIMINAIRE  = 215900   # 17pt  — AVANT-PROPOS, PROLOGUE…
SZ_SOUS_TITRE = 165100   # 13pt  — sous-titre
SZ_AUTEUR     = 190500   # 15pt  — auteur
SZ_CHAPITRE   = 177800   # 14pt  — Heading 2
SZ_CORPS      = 139700   # 11pt  — Normal
SZ_CITATION   = 133350   # 10.5pt
SZ_ATTRIBUTION= 120650   # 9.5pt

# ── Retraits (EMU) ────────────────────────────────────────────
IND_CIT_LEFT  = 432000   # 1.2cm
IND_CIT_RIGHT = 288000   # 0.8cm
IND_FIRSTLINE = 180000   # 0.5cm

# ── Styles à ne pas modifier ──────────────────────────────────
SKIP_STYLES = {
    'Heading 1', 'Heading 2', 'Heading 3',
    'Title', 'Subtitle',
    'TOC 1', 'TOC 2', 'TOC 3',
    'Caption', 'List Paragraph',
}

KEYWORDS_PARTIE    = ['PARTIE', 'ÉPILOGUE', 'EPILOGUE', 'PROLOGUE',
                      'INTRODUCTION', 'AVANT-PROPOS', 'CONCLUSION']
KEYWORDS_LIMINAIRE = ['AVANT-PROPOS', 'AVANT PROPOS', 'SOMMAIRE',
                      'PROLOGUE', 'INTRODUCTION', 'ÉPILOGUE', 'EPILOGUE']


# ─────────────────────────────────────────────────────────────
# Utilitaires XML
# ─────────────────────────────────────────────────────────────

def get_or_create_pPr(p):
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p._p.insert(0, pPr)
    return pPr


def set_spacing_xml(p, before_twips=None, after_twips=None):
    pPr = get_or_create_pPr(p)
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing')
        pPr.append(sp)
    if before_twips is not None:
        sp.set(qn('w:before'), str(before_twips))
    if after_twips is not None:
        sp.set(qn('w:after'), str(after_twips))


def add_border_left(p, sz='24', space='12', color='7B1E1E'):
    pPr = get_or_create_pPr(p)
    old = pPr.find(qn('w:pBdr'))
    if old is not None:
        pPr.remove(old)
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'),   'single')
    left.set(qn('w:sz'),    sz)
    left.set(qn('w:space'), space)
    left.set(qn('w:color'), color)
    pBdr.append(left)
    pPr.append(pBdr)


def remove_border(p):
    pPr = p._p.find(qn('w:pPr'))
    if pPr is not None:
        old = pPr.find(qn('w:pBdr'))
        if old is not None:
            pPr.remove(old)


def add_border_bottom(p, sz='8', space='8', color=GOLD):
    pPr = get_or_create_pPr(p)
    if pPr.find(qn('w:pBdr')) is not None:
        return
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    for k, v in [('w:val','single'), ('w:sz', sz),
                 ('w:space', space), ('w:color', color)]:
        bot.set(qn(k), v)
    pBdr.append(bot)
    pPr.append(pBdr)


def set_widow_control(p):
    pPr = get_or_create_pPr(p)
    wc = pPr.find(qn('w:widowControl'))
    if wc is None:
        wc = OxmlElement('w:widowControl')
        pPr.insert(0, wc)
    wc.set(qn('w:val'), '1')


# ─────────────────────────────────────────────────────────────
# MODULE 1 — Césure
# ─────────────────────────────────────────────────────────────

def apply_hyphenation(doc):
    s = doc.settings.element
    for tag, val in [('w:autoHyphenation', '1'),
                     ('w:hyphenationZone',  '720'),
                     ('w:doNotHyphenateCaps', '1')]:
        el = s.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            s.insert(0, el)
        el.set(qn('w:val'), val)
    return 1


# ─────────────────────────────────────────────────────────────
# MODULE 2 — Titres de parties (Heading 1)
# ─────────────────────────────────────────────────────────────

def apply_heading1(doc):
    count = 0
    for p in doc.paragraphs:
        if p.style.name != 'Heading 1':
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.page_break_before = True
        pf.keep_with_next    = True
        pf.space_before      = Pt(36)
        pf.space_after       = Pt(12)
        for run in p.runs:
            run.font.name      = 'Garamond'
            run.font.size      = SZ_PARTIE
            run.font.color.rgb = NAVY
            run.font.bold      = True
        add_border_bottom(p)
        count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 3 — Titres de chapitres (Heading 2)
# ─────────────────────────────────────────────────────────────

def apply_heading2(doc):
    count = 0
    for p in doc.paragraphs:
        if p.style.name != 'Heading 2':
            continue
        pf = p.paragraph_format
        pf.page_break_before = True
        pf.keep_with_next    = True
        pf.space_before      = Pt(0)
        pf.space_after       = Pt(12)
        for run in p.runs:
            run.font.name      = 'Garamond'
            run.font.size      = SZ_CHAPITRE
            run.font.color.rgb = NAVY
            run.font.bold      = False
        count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 4 — Sections liminaires
# ─────────────────────────────────────────────────────────────

def apply_liminaires(doc):
    paras = doc.paragraphs
    count = 0
    for i, p in enumerate(paras):
        txt = p.text.strip()
        if not txt or p.style.name in ('Heading 1', 'Heading 2'):
            continue
        if 47 <= i <= 155:   # intérieur sommaire → ignorer
            continue
        upper = txt.upper()
        if any(kw in upper for kw in KEYWORDS_LIMINAIRE) and len(txt) < 80:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.page_break_before = True
            pf.space_before      = Pt(24)
            pf.space_after       = Pt(18)
            pf.keep_with_next    = True
            pf.first_line_indent = 0
            for run in p.runs:
                run.font.name      = 'Garamond'
                run.font.size      = SZ_LIMINAIRE
                run.font.color.rgb = NAVY
                run.font.bold      = True
            add_border_bottom(p)
            count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 5 — Page de titre
# ─────────────────────────────────────────────────────────────

def apply_titre(doc):
    paras = doc.paragraphs
    TITRE_MAP = {
        'Alliance des États du Sahel':        (SZ_CORPS,      False, True,  NAVY,     Pt(48), Pt(6)),
        'LA PAIX,':                            (SZ_TITRE,      True,  False, NAVY,     Pt(24), Pt(0)),
        "ON PEUT L'ÉVITER":                    (SZ_TITRE,      True,  False, NAVY,     Pt(0),  Pt(18)),
        'Géopolitique de la résistance':       (SZ_SOUS_TITRE, False, True,  NAVY,     Pt(6),  Pt(60)),
        'BEN':                                 (SZ_AUTEUR,     True,  False, NAVY,     Pt(60), Pt(4)),
        '— 2026 —':                            (SZ_CORPS,      False, False, BORDEAUX, Pt(2),  Pt(2)),
    }
    count = 0
    for p in paras[:30]:
        txt = p.text.strip()
        for key, (size, bold, italic, color, sp_b, sp_a) in TITRE_MAP.items():
            if txt.startswith(key):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.space_before      = sp_b
                pf.space_after       = sp_a
                pf.first_line_indent = 0
                for run in p.runs:
                    run.font.name      = 'Garamond'
                    run.font.size      = size
                    run.font.bold      = bold
                    run.font.italic    = italic
                    run.font.color.rgb = color
                count += 1
                break
    # Bordure gold sous le sous-titre
    for p in paras[:30]:
        if p.text.strip().startswith('Géopolitique'):
            add_border_bottom(p)
    # Épigraphes sur page séparée
    for p in paras[20:45]:
        if p.text.strip().startswith('«'):
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(48)
            break
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 6 — Citations «...»
# ─────────────────────────────────────────────────────────────

def apply_citations(doc):
    count = 0
    for p in doc.paragraphs:
        if not p.text.strip().startswith('«'):
            continue
        for run in p.runs:
            run.font.size   = SZ_CITATION
            run.font.italic = True
        pf = p.paragraph_format
        pf.space_before      = Pt(8)
        pf.space_after       = Pt(2)
        pf.left_indent       = IND_CIT_LEFT
        pf.right_indent      = IND_CIT_RIGHT
        pf.first_line_indent = 0
        add_border_left(p)
        count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 7 — Attributions — Auteur
# ─────────────────────────────────────────────────────────────

def apply_attributions(doc):
    count = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not (txt.startswith('—') or txt.startswith('–')):
            continue
        if p.style.name in ('Heading 1', 'Heading 2'):
            continue
        for run in p.runs:
            run.font.size   = SZ_ATTRIBUTION
            run.font.italic = True
        pf = p.paragraph_format
        pf.space_before      = Pt(0)
        pf.space_after       = Pt(10)
        pf.left_indent       = IND_CIT_LEFT
        pf.right_indent      = IND_CIT_RIGHT
        pf.first_line_indent = 0
        remove_border(p)
        count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 8 — Étoiles séparatrices
# ─────────────────────────────────────────────────────────────

def apply_etoiles(doc):
    count = 0
    for p in doc.paragraphs:
        if '★' not in p.text:
            continue
        for run in p.runs:
            run.text = ''
        if not p.runs:
            p.add_run('★ ★ ★')
        else:
            p.runs[0].text = '★ ★ ★'
        for run in p.runs:
            run.font.color.rgb = BORDEAUX
            run.font.bold      = False
            run.font.italic    = False
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before      = Pt(10)
        pf.space_after       = Pt(10)
        pf.first_line_indent = 0
        count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 9 — Sommaire
# ─────────────────────────────────────────────────────────────

def apply_sommaire(doc):
    paras = doc.paragraphs
    n_partie = n_chap = 0
    for i, p in enumerate(paras):
        if i < 47 or i > 155:
            continue
        txt = p.text.strip()
        if not txt:
            continue
        upper = txt.upper()
        pf = p.paragraph_format
        if any(kw in upper for kw in KEYWORDS_PARTIE):
            pf.left_indent       = 0
            pf.space_before      = Pt(12)
            pf.space_after       = Pt(4)
            pf.first_line_indent = 0
            for run in p.runs:
                run.font.bold = True
            n_partie += 1
        elif 'Chapitre' in txt or 'CHAPITRE' in upper:
            pf.left_indent       = Cm(1.5)
            pf.space_before      = Pt(3)
            pf.space_after       = 0
            pf.first_line_indent = 0
            for run in p.runs:
                run.font.bold = False
            n_chap += 1
    return n_partie, n_chap


# ─────────────────────────────────────────────────────────────
# MODULE 10 — Tableaux
# ─────────────────────────────────────────────────────────────

def apply_tableaux(doc):
    for table in doc.tables:
        # En-tête : sp_before = 14pt (280 twips)
        header = table.rows[0].cells[0].paragraphs[0]
        set_spacing_xml(header, before_twips=280)
        # Normaliser fond en-tête sur #1A1A2E
        tc   = table.rows[0].cells[0]._tc
        tcPr = tc.find(qn('w:tcPr'))
        if tcPr is not None:
            shd = tcPr.find(qn('w:shd'))
            if shd is not None:
                shd.set(qn('w:fill'), '1A1A2E')
        # Dernière ligne : sp_after = 8pt (160 twips)
        last = table.rows[-1].cells[0].paragraphs[-1]
        set_spacing_xml(last, after_twips=160)
    return len(doc.tables)


# ─────────────────────────────────────────────────────────────
# MODULE 11 — Corps de texte
# ─────────────────────────────────────────────────────────────

def apply_corps(doc):
    count = 0
    for p in doc.paragraphs:
        if p.style.name in SKIP_STYLES:
            continue
        txt = p.text.strip()
        if not txt or len(txt) < 20:
            continue
        if txt.startswith(('«', '—', '–')) or '★' in txt:
            continue
        if p.style.name in ('Normal', 'Body Text', '') \
                or p.style.name.startswith('Normal'):
            p.alignment          = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.line_spacing      = 1.3
            pf.widow_control     = True
            pf.space_after       = Pt(6)
            set_widow_control(p)
            count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MODULE 12 — Retrait première ligne
# ─────────────────────────────────────────────────────────────

def apply_retrait(doc):
    paras = doc.paragraphs
    count = 0
    skip_next = False
    for i, p in enumerate(paras):
        if p.style.name in ('Heading 1', 'Heading 2'):
            skip_next = True
            continue
        txt = p.text.strip()
        if not txt:
            continue
        if 47 <= i <= 155:
            skip_next = False
            continue
        if p.style.name in SKIP_STYLES:
            skip_next = False
            continue
        if txt.startswith(('«', '—', '–')) or '★' in txt:
            skip_next = False
            continue
        if any(kw in txt.upper() for kw in KEYWORDS_LIMINAIRE) and len(txt) < 80:
            skip_next = False
            continue
        if len(txt) < 40:
            skip_next = False
            continue
        # 1er § après un titre → pas de retrait
        if skip_next:
            p.paragraph_format.first_line_indent = 0
            skip_next = False
            continue
        if p.style.name in ('Normal', 'Body Text', '') \
                or p.style.name.startswith('Normal'):
            p.paragraph_format.first_line_indent = IND_FIRSTLINE
            count += 1
    return count


# ─────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='Formateur KDP professionnel')
    parser.add_argument('--input',  default='LA_PAIX_ON_PEUT_L_EVITER_CORRIGE.docx')
    parser.add_argument('--output', default=None)
    args = parser.parse_args()

    output = args.output or args.input

    print(f"Chargement : {args.input}")
    doc = Document(args.input)
    print(f"  {len(doc.paragraphs)} paragraphes, {len(doc.tables)} tableaux\n")

    print("MODULE 1  — Césure")
    apply_hyphenation(doc)
    print("  ✓ autoHyphenation activée\n")

    print("MODULE 2  — Titres de parties (Heading 1)")
    n = apply_heading1(doc)
    print(f"  ✓ {n} PARTIES\n")

    print("MODULE 3  — Titres de chapitres (Heading 2)")
    n = apply_heading2(doc)
    print(f"  ✓ {n} chapitres\n")

    print("MODULE 4  — Sections liminaires")
    n = apply_liminaires(doc)
    print(f"  ✓ {n} sections\n")

    print("MODULE 5  — Page de titre")
    n = apply_titre(doc)
    print(f"  ✓ {n} éléments\n")

    print("MODULE 6  — Citations «...»")
    n = apply_citations(doc)
    print(f"  ✓ {n} citations\n")

    print("MODULE 7  — Attributions")
    n = apply_attributions(doc)
    print(f"  ✓ {n} attributions\n")

    print("MODULE 8  — Étoiles ★ ★ ★")
    n = apply_etoiles(doc)
    print(f"  ✓ {n} étoiles\n")

    print("MODULE 9  — Sommaire")
    n_p, n_c = apply_sommaire(doc)
    print(f"  ✓ {n_p} PARTIE, {n_c} chapitres\n")

    print("MODULE 10 — Tableaux")
    n = apply_tableaux(doc)
    print(f"  ✓ {n} tableaux\n")

    print("MODULE 11 — Corps de texte")
    n = apply_corps(doc)
    print(f"  ✓ {n} paragraphes\n")

    print("MODULE 12 — Retrait première ligne")
    n = apply_retrait(doc)
    print(f"  ✓ {n} paragraphes\n")

    doc.save(output)
    print(f"{'─'*50}")
    print(f"Fichier sauvegardé : {output}")
    print(f"{'─'*50}")


if __name__ == '__main__':
    main()
